from datetime import datetime

import openpyxl
import pytest
from openpyxl.styles import PatternFill

HEADER = ["Tanggal", "ISSUE", "Aplikasi", "Pelapor", "Media Laporan", "Penerima Laporan",
          "Pencatat Issue", "Waktu Lapor", "Status", "Pic", "Waktu Mulai", "Tanggal Finish",
          "Waktu Selesai", "Note", "Aplikasi (IP)", "Dokumentasi"]

GREEN = PatternFill("solid", fgColor="FFB7E1CD")
LIGHT_GREEN = PatternFill("solid", fgColor="FFD9EAD3")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")


def _row(tgl, app, pelapor, status="", pic="", note=""):
    return [tgl, "deskripsi issue", app, pelapor, "Whatsapp", "RCO Official", "Dana",
            "09:00", status, pic, "", "", "", note, "172.18.141.71", ""]


@pytest.fixture
def workbook_path(tmp_path):
    path = tmp_path / "handover.xlsx"
    make_sample_workbook(path)
    return path


def make_sample_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issue Oktober 2026"
    ws.append([f"Column {i}" for i in range(1, 17)])   # baris 1: tabel otomatis
    ws.append(HEADER)                                   # baris 2: header
    rows = [
        _row(datetime(2026, 10, 1), "MPAYS", "1 ~Aldi Pratama"),                       # 3 pending Tim 1
        _row(datetime(2026, 10, 1), "QRIS", "2 ~Putri", status="oncek", pic="Hendrik"),  # 4 progress Tim 2
        _row(datetime(2026, 10, 2), "JCB", "3 temuan", status="Done", pic="Elsa"),      # 5 done status
        _row(datetime(2026, 10, 2), "BRIMO", "IT Service Desk"),                        # 6 hijau -> done
        _row("05/10/2026", "Swift MT940", "Tim 1 - Budi", pic="select * from x"),        # 7 PIC SQL -> pending
        _row(datetime(2026, 10, 7), "EDC", 2.0, status="tbd"),                          # 8 progress tanpa PIC
        _row(datetime(2026, 10, 9), "BI-FAST", "62 811-1122-993 ~Aldi"),                # 9 pending ALL TEAM
        _row(datetime(2026, 10, 10), "QRIS", "1 ~X", status="30/12/1899", pic="-"),     # 10 pending
        _row(datetime(2026, 10, 11), "QRIS", "3 ~Y", pic="done"),                       # 11 done via pic
        _row(datetime(2026, 10, 12), "MPAYS", "1 ~Z"),                                  # 12 light green -> done
        _row(datetime(2026, 10, 13), "MPAYS", "2 ~Q", pic="Puji"),                      # 13 kuning fill, progress
        _row(datetime(2026, 9, 30), "X", "1 ~W"),                                       # 14 tanggal di luar Oktober
    ]
    for r in rows:
        ws.append(r)
    for c in range(1, 17):
        ws.cell(row=6, column=c).fill = GREEN
    ws.cell(row=12, column=1).fill = LIGHT_GREEN
    for c in range(1, 17):
        ws.cell(row=13, column=c).fill = YELLOW

    # Tab tanpa header: data langsung dari baris 3
    ws2 = wb.create_sheet("Issue September 2026")
    ws2.append([None] * 16)
    ws2.append([None] * 16)
    ws2.append(_row(datetime(2026, 9, 28), "QRIS", "2 ~A"))                              # 3 pending Tim 2
    ws2.append(_row(datetime(2026, 9, 29), "QRIS", "3 ~B", status="on progress", pic="Azhari"))  # 4 progress

    wb.save(path)


if __name__ == "__main__":
    import sys

    make_sample_workbook(sys.argv[1])
    print(f"Contoh tersimpan di {sys.argv[1]}")
